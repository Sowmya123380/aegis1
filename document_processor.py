
from io import BytesIO
from pathlib import Path
from difflib import SequenceMatcher
import re

import pandas as pd
from docx import Document
from pypdf import PdfReader


def extract_text(filename, file_bytes):
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        pages = []

        for number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[Page {number}]\n{page_text}")

        text = "\n\n".join(pages)

    elif ext == ".docx":
        doc = Document(BytesIO(file_bytes))
        parts = [
            p.text for p in doc.paragraphs if p.text.strip()
        ]

        for table in doc.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

        text = "\n".join(parts)

    elif ext == ".csv":
        df = pd.read_csv(BytesIO(file_bytes))
        text = df.to_string(index=False)

    else:
        raise ValueError(
            "Supported formats: PDF, DOCX and CSV. "
            "Legacy .doc files are not supported."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. This may be a scanned PDF. "
            "OCR is required to read image-based pages."
        )

    return text


def extract_key_information(text):
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip() and not line.strip().startswith("[Page ")
    ]

    emails = sorted(set(re.findall(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )))

    dates = sorted(set(re.findall(
        r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|"
        r"\d{4}-\d{2}-\d{2}|"
        r"\d{1,2}\s+[A-Za-z]+\s+\d{4})\b",
        text
    )))

    amounts = sorted(set(re.findall(
        r"(?:₹|Rs\.?\s*|\$\s*|USD\s*|INR\s*)"
        r"\d[\d,]*(?:\.\d{1,2})?",
        text,
        flags=re.IGNORECASE
    )))

    return {
        "word_count": len(text.split()),
        "key_lines": lines[:12],
        "emails": emails,
        "dates": dates,
        "amounts": amounts,
    }


def summarize_locally(text, max_sentences=6):
    # A simple extractive summary; does not call an AI API.
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)

    cleaned = [
        sentence.strip()
        for sentence in sentences
        if len(sentence.strip()) > 25
        and not sentence.strip().startswith("[Page ")
    ]

    if not cleaned:
        return text[:1500]

    # Select sentences from across the document.
    count = min(max_sentences, len(cleaned))
    if count == 1:
        return cleaned[0]

    indices = [
        round(i * (len(cleaned) - 1) / (count - 1))
        for i in range(count)
    ]

    return "\n\n".join(cleaned[i] for i in indices)


def compare_documents(text_a, text_b):
    lines_a = [
        line.strip() for line in text_a.splitlines()
        if line.strip() and not line.strip().startswith("[Page ")
    ]
    lines_b = [
        line.strip() for line in text_b.splitlines()
        if line.strip() and not line.strip().startswith("[Page ")
    ]

    normalized_a = [line.casefold() for line in lines_a]
    normalized_b = [line.casefold() for line in lines_b]

    similarity = SequenceMatcher(
        None, text_a.casefold(), text_b.casefold()
    ).ratio() * 100

    set_a = set(normalized_a)
    set_b = set(normalized_b)

    return {
        "similarity_percent": round(similarity, 2),
        "only_in_document_1": [
            line for line in lines_a if line.casefold() not in set_b
        ],
        "only_in_document_2": [
            line for line in lines_b if line.casefold() not in set_a
        ],
        "common_lines": [
            line for line in lines_a if line.casefold() in set_b
        ],
    }